# wb_uploader v2.0
# G2G code
#
# Один запуск = весь путь от готовой таблицы до карточек с фото на Wildberries:
#   0. проверка всего локального ДО отправки (таблица, папки с фото, код ТН ВЭД) — если что-то
#      не так, скрипт останавливается, ничего не создав;
#   1. правка слишком длинных значений прямо в консоли (в таблицу не пишется, уходит только в WB);
#   2. создание карточек через API, ожидание их появления; не создавшиеся карточки
#      пробуются ещё RETRIES раз, после чего скрипт останавливается и показывает причину;
#   3. проставление кода ТН ВЭД в каждой карточке (в шаблоне WB он не загружается);
#   4. загрузка фото.
#
# Порядок работы:
#   1) mockuper.jsx  2) excel_filler.py  3) правите Постеры.xlsx (цвета и т.п.)  4) этот скрипт
#   5) остатки — вручную.
#
# Ожидаемое расположение (скрипт лежит в <корень>/scripts/):
#   <корень>/Постеры.xlsx
#   <корень>/mockups/{8.jpg, cover.jpg}                                   — общие для всех карточек
#   <корень>/output/<Исполнитель>/<исходник>/{1,2,3,4,5-3|5-4|5-5,6,7}.jpg
#
# Повторный запуск безопасен: карточки, которые уже есть в WB, заново не создаются,
# ТН ВЭД ставится только там, где он отличается, фото пропускаются, если уже загружены.
#
# Группа: строки таблицы с одинаковым значением в столбце "Группа" объединяются в одну
# карточку WB (как при ручной загрузке шаблона). Чтобы каждая строка была отдельной
# карточкой — поставьте MERGE_BY_GROUP = False.

import json
import re
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

# ------------------- КОНФИГ / ТОКЕН -------------------
API_TOKEN = ""
BASE_URL = "https://content-api.wildberries.ru"
HEADERS = {}

# Колбэки используются будущим GUI. При обычном запуске из Terminal
# сообщения и запросы работают почти так же, как в старой версии.
LOG_CALLBACK = None
PROMPT_CALLBACK = None
PROGRESS_CALLBACK = None


def _log(*args, **kwargs):
    """Сообщение одновременно выводится в Terminal и может уйти в GUI."""
    import builtins
    builtins.print(*args, **kwargs)
    if LOG_CALLBACK:
        try:
            message = " ".join(str(a) for a in args)
            LOG_CALLBACK(message)
        except Exception:
            pass


def _progress(current, total, message=""):
    """Передаёт прогресс GUI в диапазоне 0..100."""
    if PROGRESS_CALLBACK:
        try:
            value = 0 if total <= 0 else max(0, min(100, int(current / total * 100)))
            PROGRESS_CALLBACK(value, message)
        except Exception:
            pass


def _prompt(message, default=None):
    """Запрос значения: GUI может предоставить свой диалог."""
    if PROMPT_CALLBACK:
        value = PROMPT_CALLBACK(message, default)
        if value is None:
            return "" if default is None else str(default)
        return str(value)
    import builtins
    value = builtins.input(message)
    return value.strip()


def _confirm(message, default=False):
    """Да/нет запрос. Для GUI prompt_callback возвращает текстовый ответ."""
    answer = _prompt(message, "y" if default else "")
    if not answer and default:
        return True
    return answer.strip().lower() in ("y", "yes", "д", "да")


def load_config(config_path=None):
    """Читает токен WB из config.json в папке scripts/.

    Формат:
    {
        "wildberries": {
            "api_token": "..."
        }
    }
    """
    path = Path(config_path) if config_path else CONFIG_FILE
    if not path.exists():
        raise RuntimeError(
            f"Не найден config.json: {path}\n"
            "Создайте его и добавьте: {\"wildberries\": {\"api_token\": \"ВАШ_ТОКЕН\"}}"
        )

    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception as e:
        raise RuntimeError(f"Не удалось прочитать config.json: {e}")

    token = ((data.get("wildberries") or {}).get("api_token") or "").strip()
    if not token:
        raise RuntimeError("В config.json не заполнен wildberries.api_token")

    return token


# config.json лежит в scripts/ рядом со скриптами:
# {
#   "wildberries": {
#     "api_token": "ВАШ_ТОКЕН_WB"
#   }
# }

# ------------------- РАСПОЛОЖЕНИЕ ПРОЕКТА -------------------
SCRIPT_DIR = Path(__file__).resolve().parent
ROOT_DIR = SCRIPT_DIR.parent
CONFIG_FILE = SCRIPT_DIR / "config.json"

EXCEL_FILE = ROOT_DIR / "Постеры.xlsx"
OUTPUT_FOLDER = ROOT_DIR / "output"
MOCKUPS_FOLDER = ROOT_DIR / "mockups"
# Запоминает, как вы переименовали слишком длинные артикулы. Нужен только для повторного
# запуска после сбоя (чтобы не создать дубли под другим артикулом). После полностью
# успешного прогона файл удаляется.
RENAMES_FILE = SCRIPT_DIR / "wb_uploader_renames.json"

# ------------------- НАСТРОЙКИ -------------------
SHEET_NAME = "Товары"
ARTICLES_COLUMN = "Артикул продавца"
STOP_WORDS = ("номер", "идентификатор")  # строки-подсказки в шапке шаблона WB

TNVED_COLUMN = "Код ТН ВЭД"
TNVED_CHARC_ID_FALLBACK = 15001135  # запасной ID характеристики "ТНВЭД"

MERGE_BY_GROUP = True
MAX_VARIANTS_PER_CARD = 30

# Как в excel_filler: значение такой длины и больше считается слишком длинным (красная ячейка)
MAX_TITLE_LEN = 60
MAX_ARTICLE_LEN = 72

# Порядок "своих" фото исходника и общие фото (как в wb_photo_uploader v1.7)
PHOTO_ORDER = ["1", "2", "3", "4", "5-x", "6", "7"]
STATIC_PHOTOS = ["8", "cover"]
FORCE_REUPLOAD = False

# Столбцы таблицы, которые НЕ являются характеристиками предмета
NON_CHARC_COLUMNS = [
    "Группа", "Артикул продавца", "Артикул WB", "Наименование", "Категория продавца",
    "Бренд", "Описание", "Фото", "Видео", "Баркоды", "Цена", "Артикул OZON",
    "Вес с упаковкой (кг)", "Длина упаковки", "Ширина упаковки", "Высота упаковки", "КИЗ", "18+",
    "Только для ИП и юрлиц", "Подтверждаю, что товар промаркирован",
]

RETRIES = 3               # сколько раз повторно пытаться создать не создавшиеся карточки
CREATE_BATCH = 50         # карточек (групп) за один запрос создания
UPDATE_BATCH = 100
UPLOAD_PAUSE = 7          # сек между запросами создания/обновления (лимит WB — 10 в минуту)
WAIT_CREATE_TIMEOUT = 600  # сек ожидания создания одной попытки
WAIT_POLL = 20
WAIT_TNVED_TIMEOUT = 240


# ------------------- ВСПОМОГАТЕЛЬНОЕ -------------------
def norm(text):
    return re.sub(r"[^0-9a-zа-яё]", "", str(text).lower())


def chunks(items, size):
    for i in range(0, len(items), size):
        yield items[i:i + size]


def parse_ts(value):
    try:
        ts = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return ts if ts.tzinfo else ts.replace(tzinfo=timezone.utc)
    except (ValueError, TypeError):
        return None


def api(method, url, json_body=None, params=None, retries=5):
    """Запрос к WB с повторами на 429/5xx и сетевые ошибки."""
    for attempt in range(retries):
        try:
            r = requests.request(method, url, headers=HEADERS, json=json_body, params=params, timeout=90)
        except requests.RequestException as e:
            _log(f"!  Сетевая ошибка ({e}), повтор...")
            time.sleep(3 * (attempt + 1))
            continue
        if r.status_code == 401:
            raise SystemExit("x WB вернул 401: токен неверный или просрочен.")
        if r.status_code == 429 or r.status_code >= 500:
            try:
                wait = float(r.headers.get("X-Ratelimit-Retry", 0))
            except ValueError:
                wait = 0
            time.sleep(wait or 5 * (attempt + 1))
            continue
        return r
    raise RuntimeError(f"{method} {url}: нет успешного ответа после {retries} попыток")


def json_or_empty(r):
    try:
        return r.json()
    except ValueError:
        return {}


# ------------------- ТАБЛИЦА -------------------
def clean_value(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def read_table():
    """Строки таблицы: список словарей {название столбца: значение} для валидных артикулов."""
    try:
        wb = load_workbook(EXCEL_FILE, data_only=True)
    except PermissionError:
        raise SystemExit("x Не удалось открыть Постеры.xlsx — закройте файл в Excel и запустите снова.")
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        header_row = None
        for r in range(1, min(ws.max_row, 30) + 1):
            for c in ws[r]:
                if isinstance(c.value, str) and c.value.strip().lower() == ARTICLES_COLUMN.lower():
                    header_row = r
                    break
            if header_row:
                break
        if not header_row:
            raise SystemExit(f"x В таблице не найден столбец '{ARTICLES_COLUMN}'")

        headers = {c.column: c.value.strip() for c in ws[header_row] if isinstance(c.value, str) and c.value.strip()}
        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            row = {name: clean_value(ws.cell(row=r, column=col).value) for col, name in headers.items()}
            article = row.get(ARTICLES_COLUMN)
            if not article:
                continue
            article = str(article).strip()
            if any(w in article.lower() for w in STOP_WORDS):
                continue
            row[ARTICLES_COLUMN] = article
            row["_orig_article"] = article  # по нему ищется папка с фото, даже если артикул переименован
            row["_row"] = r
            rows.append(row)
        return rows
    finally:
        wb.close()


def make_title(row):
    """Итоговое 'Наименование', которое отправляется в WB.
    Если оно задано в таблице — используем его как есть; иначе собираем
    из 'Наименование постера' + 'А3' + 'Цвет'."""
    title = row.get("Наименование")
    if title:
        return str(title).strip()
    poster = row.get("Наименование постера") or ""
    color = row.get("Цвет") or ""
    return f"{poster} А3 {color}".strip()


def tnved_of_row(row):
    v = row.get(TNVED_COLUMN)
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


# ------------------- ФОТО НА ДИСКЕ -------------------
def get_artist_name(artist=None):
    name = str(artist).strip() if artist is not None else _prompt("Введите имя исполнителя: ").strip()
    if not name:
        raise SystemExit("Имя исполнителя не введено — остановка.")
    return name


def find_artist_output_folder(artist_name):
    if not OUTPUT_FOLDER.exists():
        return None
    target = artist_name.strip().lower()
    for sub in OUTPUT_FOLDER.iterdir():
        if sub.is_dir() and sub.name.strip().lower() == target:
            return sub
    return None


def build_folder_index(artist_folder, artist_name):
    """article -> папка с фото: '<исходник в нижнем регистре> <исполнитель в нижнем регистре>'."""
    artist_lower = artist_name.strip().lower()
    return {f"{sub.name.strip().lower()} {artist_lower}": sub
            for sub in artist_folder.iterdir() if sub.is_dir()}


def collect_photos(source_folder):
    photos = []
    for key in PHOTO_ORDER:
        if key == "5-x":
            for variant in ("5-3", "5-4", "5-5"):
                candidate = source_folder / f"{variant}.jpg"
                if candidate.exists():
                    photos.append(candidate)
                    break
            continue
        candidate = source_folder / f"{key}.jpg"
        if candidate.exists():
            photos.append(candidate)
    return photos


def collect_static_photos():
    result = []
    for stem in STATIC_PHOTOS:
        candidate = MOCKUPS_FOLDER / f"{stem}.jpg"
        if candidate.exists():
            result.append(candidate)
        else:
            _log(f"!  Не найден общий файл '{candidate}' — будет пропущен для всех карточек.")
    return result


# ------------------- ПРОВЕРКА ДО ОТПРАВКИ -------------------
def validate_group_limits(rows):
    """Проверяет лимит WB: в одной группе должно быть не больше 30 строк."""
    counts = {}
    for r in rows:
        value = r.get("Группа")
        if value is None:
            continue
        key = str(value).strip()
        if not key:
            continue
        counts[key] = counts.get(key, 0) + 1

    exceeded = [(group, count) for group, count in counts.items()
                if count > MAX_VARIANTS_PER_CARD]
    if exceeded:
        _log("x В столбце 'Группа' превышен лимит WB — в одной группе не может быть больше "
              f"{MAX_VARIANTS_PER_CARD} строк.")
        for group, count in exceeded:
            _log(f"   Группа '{group}': {count} строк.")
        _log("   Исправьте значения в столбце 'Группа' и запустите скрипт снова.")
        raise SystemExit(1)


def validate_local(rows, folder_index):
    problems = []
    seen = {}
    for r in rows:
        article = r[ARTICLES_COLUMN]
        low = article.lower()
        if low in seen:
            problems.append(f"{article}: артикул повторяется (строки {seen[low]} и {r['_row']})")
        seen[low] = r["_row"]

        if not r.get("Цвет"):
            problems.append(f"{article}: не заполнен 'Цвет'")
        title = make_title(r)
        if not title or '""' in title:
            problems.append(f"{article}: не заполнено 'Наименование'")

        code = tnved_of_row(r)
        if not code:
            problems.append(f"{article}: пустой '{TNVED_COLUMN}'")
        elif not re.fullmatch(r"\d{10}", code):
            problems.append(f"{article}: '{TNVED_COLUMN}' = {code} — ожидается 10 цифр")

        folder = folder_index.get(article.lower())
        if not folder:
            problems.append(f"{article}: не найдена папка с фото в output/<исполнитель>/")
        elif not collect_photos(folder):
            problems.append(f"{article}: в папке '{folder.name}' нет ни одного фото из списка")
    if problems:
        _log("x Найдены проблемы, ничего не отправлено в WB:")
        for line in problems:
            _log("   " + line)
        raise SystemExit(1)


# ------------------- WILDBERRIES: СПРАВОЧНИКИ -------------------
def find_subject_id(name):
    r = api("GET", BASE_URL + "/content/v2/object/all", params={"name": name, "locale": "ru", "limit": 1000})
    r.raise_for_status()
    items = r.json().get("data") or []
    exact = [i for i in items if str(i.get("subjectName", "")).strip().lower() == name.strip().lower()]
    if not exact:
        _log(f"x Предмет '{name}' не найден. Похожие: {[i.get('subjectName') for i in items[:10]]}")
        raise SystemExit(1)
    if len(exact) > 1:
        _log(f"!  Найдено несколько предметов '{name}', беру первый (subjectID {exact[0]['subjectID']}).")
    return exact[0]["subjectID"]


def load_charcs(subject_id):
    r = api("GET", f"{BASE_URL}/content/v2/object/charcs/{subject_id}", params={"locale": "ru"})
    r.raise_for_status()
    return {norm(c["name"]): c for c in (r.json().get("data") or [])}


def load_color_directory():
    """Загружает актуальный справочник допустимых значений характеристики 'Цвет' WB."""
    r = api("GET", BASE_URL + "/content/v2/directory/colors", params={"locale": "ru"})
    r.raise_for_status()
    data = r.json().get("data") or []
    # Ключ — нормализованное название, значение — оригинальное название WB.
    return {str(item.get("name", "")).strip().casefold(): str(item.get("name", "")).strip()
            for item in data if str(item.get("name", "")).strip()}


def check_color_values(rows):
    """Проверяет 'Цвет' и просит исправить недопустимое значение прямо в консоли.

    Изменение делается только в памяти — исходный Excel не меняется.
    В WB уйдёт уже исправленное значение.
    """
    allowed = load_color_directory()
    if not allowed:
        raise SystemExit("x WB не вернул справочник цветов — создание карточек остановлено для безопасности.")

    import difflib
    allowed_names = list(allowed.values())

    for row in rows:
        raw = row.get("Цвет")
        if raw is None or not str(raw).strip():
            continue  # Пустое значение уже ловит validate_local().

        while True:
            values = [part.strip() for part in str(row.get("Цвет") or "").split(";") if part.strip()]
            invalid = [value for value in values if value.casefold() not in allowed]
            if not invalid:
                # Приводим корректные значения к написанию WB.
                row["Цвет"] = "; ".join(allowed[value.casefold()] for value in values)
                break

            _log("\nЦвет не найден в допустимых значениях WB:")
            _log(f"   строка {row['_row']}, артикул '{row[ARTICLES_COLUMN]}'")
            _log(f"   Сейчас: {row.get('Цвет')}")
            _log(f"   Недопустимо: {', '.join(repr(v) for v in invalid)}")

            for value in invalid:
                suggestions = difflib.get_close_matches(value, allowed_names, n=5, cutoff=0.35)
                if suggestions:
                    _log(f"   Возможные значения WB для '{value}': {', '.join(suggestions)}")

            def color_ok(value):
                if not value.strip():
                    return False, "Пустое значение."
                parts = [part.strip() for part in value.split(";") if part.strip()]
                bad = [part for part in parts if part.casefold() not in allowed]
                if bad:
                    return False, f"Недопустимые значения: {', '.join(bad)}"
                return True, ""

            row["Цвет"] = ask_value(
                "Новое значение 'Цвет' (несколько через ';'): ",
                color_ok,
                default=None,
            )


def find_tnved_charc(charcs):
    for key, c in charcs.items():
        if "тнвэд" in key:
            return c["charcID"], c["name"]
    _log(f"!  Характеристика ТН ВЭД не найдена среди характеристик предмета, "
          f"использую ID {TNVED_CHARC_ID_FALLBACK}.")
    return TNVED_CHARC_ID_FALLBACK, "ТНВЭД (запасной ID)"


def tnved_in_directory(subject_id, code):
    """True/False, либо None если справочник не ответил."""
    r = api("GET", BASE_URL + "/content/v2/directory/tnved",
            params={"subjectID": subject_id, "search": code, "locale": "ru"})
    if r.status_code != 200:
        _log(f"!  Не удалось проверить код {code} по справочнику ТН ВЭД (HTTP {r.status_code}).")
        return None
    codes = [str(i.get("tnved")) for i in (r.json().get("data") or [])]
    if code in codes:
        return True
    r = api("GET", BASE_URL + "/content/v2/directory/tnved",
            params={"subjectID": subject_id, "search": code[:4], "locale": "ru"})
    if r.status_code == 200:
        options = [str(i.get("tnved")) for i in (r.json().get("data") or [])]
        if options:
            _log(f"   Допустимые коды с началом {code[:4]}: {', '.join(options[:15])}")
    return False


def check_tnved_codes(rows, subject_id):
    for code in sorted({tnved_of_row(r) for r in rows}):
        result = tnved_in_directory(subject_id, code)
        if result is False:
            _log(f"!  Кода {code} нет в справочнике WB для этого предмета — WB может его отклонить.")
            if not _confirm("   Продолжить всё равно? [y/N]: "):
                raise SystemExit("Остановлено пользователем.")
        elif result:
            _log(f"   ТН ВЭД {code}: есть в справочнике WB.")


# ------------------- WILDBERRIES: СПИСОК КАРТОЧЕК -------------------
def fetch_cards(subject_id, wanted=None, since=None):
    """{vendorCode: карточка WB}. wanted — остановиться, как только найдены все артикулы;
    since — не листать дальше карточек, обновлённых раньше этого момента (список идёт от новых)."""
    result = {}
    cursor = {"limit": 100}
    wanted = set(wanted) if wanted else None
    while True:
        payload = {"settings": {
            "sort": {"ascending": False},
            "cursor": cursor,
            "filter": {"withPhoto": -1, "objectIDs": [subject_id]},
        }}
        r = api("POST", BASE_URL + "/content/v2/get/cards/list", json_body=payload, params={"locale": "ru"})
        r.raise_for_status()
        data = r.json()
        cards = data.get("cards") or []
        for card in cards:
            vendor = str(card.get("vendorCode", "")).strip()
            if vendor and card.get("nmID"):
                result[vendor] = card
        cur = data.get("cursor") or {}
        if wanted is not None and wanted.issubset(result.keys()):
            break
        if not cards or cur.get("total", 0) < 100:
            break
        if since and cards:
            last_ts = parse_ts(cards[-1].get("updatedAt"))
            if last_ts and last_ts < since:
                break
        cursor = {"limit": 100, "updatedAt": cur.get("updatedAt"), "nmID": cur.get("nmID")}
        time.sleep(0.7)
    return result


def get_creation_errors(vendors, since):
    """{vendorCode: текст ошибки} по списку несозданных карточек WB (только свежие)."""
    found = {}
    try:
        r = api("POST", BASE_URL + "/content/v2/cards/error/list",
                json_body={"cursor": {"limit": 100}, "order": {"ascending": False}})
        if r.status_code != 200:
            return found
        for item in r.json().get("data") or []:
            ts = parse_ts(item.get("updatedAt") or item.get("updateAt"))
            if ts and since and ts < since:
                continue
            text = "; ".join(str(e) for e in (item.get("errors") or [])) or "ошибка без описания"
            for v in item.get("vendorCodes") or []:
                if v in vendors:
                    found[v] = text
    except Exception:
        pass
    return found


# ------------------- ПРАВКА ДЛИННЫХ ЗНАЧЕНИЙ В КОНСОЛИ -------------------
def load_renames():
    if RENAMES_FILE.exists():
        try:
            return json.loads(RENAMES_FILE.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            pass
    return {}


def save_renames(renames):
    try:
        RENAMES_FILE.write_text(json.dumps(renames, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError as e:
        _log(f"!  Не удалось сохранить {RENAMES_FILE.name}: {e}")


def ask_value(prompt, is_valid, default=None):
    while True:
        value = _prompt(prompt, default).strip()
        if not value and default:
            value = default
        ok, message = is_valid(value)
        if ok:
            return value
        _log(f"   {message}")


def fix_lengths(todo, all_rows, renames):
    """Для создаваемых карточек: если итоговое 'Наименование' слишком длинное —
    показывает его целиком и просит новое значение. Правка живёт только в памяти
    и уходит только в WB; таблица не изменяется."""
    # --- итоговое наименование ---
    for r in todo:
        title = make_title(r)
        if len(title) < MAX_TITLE_LEN:
            continue

        _log(f"\\nНаименование слишком длинное (лимит — меньше {MAX_TITLE_LEN} симв.):")
        _log(f"   {len(title):>3} симв.: {title}")

        def title_ok(value):
            if not value:
                return False, "Пустое значение."
            if len(value) >= MAX_TITLE_LEN:
                return False, f"Всё ещё слишком длинно ({len(value)} симв.)."
            return True, ""

        new_title = ask_value(
            f"Новое 'Наименование' (до {MAX_TITLE_LEN - 1} симв.): ",
            title_ok,
        )
        r["Наименование"] = new_title

    # --- артикулы ---
    taken = {r[ARTICLES_COLUMN].lower() for r in all_rows}
    for r in todo:
        article = r[ARTICLES_COLUMN]
        if len(article) < MAX_ARTICLE_LEN:
            continue
        _log(f"\\nАртикул продавца слишком длинный ({len(article)} симв., лимит — меньше {MAX_ARTICLE_LEN}):")
        _log(f"   {article}")
        saved = renames.get(r["_orig_article"])
        if saved:
            _log(f"   В прошлый раз вы вводили: {saved}  (Enter — использовать его)")

        def article_ok(value, current=article):
            if not value:
                return False, "Пустое значение."
            if len(value) >= MAX_ARTICLE_LEN:
                return False, f"Всё ещё слишком длинно ({len(value)} симв.)."
            if value.lower() in taken and value.lower() != current.lower():
                return False, "Такой артикул уже есть в таблице."
            return True, ""

        new_article = ask_value(
            f"Новый артикул продавца (до {MAX_ARTICLE_LEN - 1} симв.): ",
            article_ok,
            default=saved,
        )
        taken.add(new_article.lower())
        r[ARTICLES_COLUMN] = new_article
        renames[r["_orig_article"]] = new_article
        save_renames(renames)


def to_charc_value(charc, raw):
    if charc.get("charcType") == 4:  # числовая характеристика
        try:
            number = float(str(raw).replace(",", "."))
        except ValueError:
            return None
        return int(number) if number.is_integer() else number
    text = str(int(raw)) if isinstance(raw, float) and raw.is_integer() else str(raw).strip()
    max_count = charc.get("maxCount", 0)
    values = [p.strip() for p in text.split(";") if p.strip()] if (max_count != 1 and ";" in text) else [text]
    if max_count and len(values) > max_count:
        values = values[:max_count]
    return values


def num(v, default=0):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return default


def build_variant(row, charcs, barcode, unmatched):
    non_charc = {norm(n) for n in NON_CHARC_COLUMNS}
    characteristics = []
    for name, raw in row.items():
        if name.startswith("_"):
            continue
        key = norm(name)
        # Вес с упаковкой теперь передаётся не как характеристика,
        # а в variants[].dimensions[].weightBrutto.
        if raw is None or key in non_charc or "тнвэд" in key or key == norm("Вес с упаковкой (кг)") or key == norm("Вес товара с упаковкой (г)"):
            continue
        charc = charcs.get(key)
        if not charc:
            unmatched.add(name)
            continue
        value = to_charc_value(charc, raw)
        if value is None or value == []:
            continue
        characteristics.append({"id": charc["charcID"], "value": value})

    weight_kg = num(row.get("Вес с упаковкой (кг)"))
    if not weight_kg:
        weight_kg = round(num(row.get("Вес товара с упаковкой (г)")) / 1000, 3)

    variant = {
        "vendorCode": row[ARTICLES_COLUMN],
        "title": make_title(row),
        "description": str(row.get("Описание") or ""),
        # WB ожидает dimensions как объект.
        # Вес с упаковкой — в килограммах и внутри dimensions.weightBrutto.
        "dimensions": {
            "length": int(round(num(row.get("Длина упаковки")))),
            "width": int(round(num(row.get("Ширина упаковки")))),
            "height": int(round(num(row.get("Высота упаковки")))),
            "weightBrutto": weight_kg,
        },
        "characteristics": characteristics,
        "sizes": [{"price": int(round(num(row.get("Цена")))), "skus": [barcode]}],
    }
    if row.get("Бренд"):
        variant["brand"] = str(row["Бренд"])
    return variant


def generate_barcodes(count):
    r = api("POST", BASE_URL + "/content/v2/barcodes", json_body={"count": count})
    r.raise_for_status()
    return r.json().get("data") or []


def make_groups(rows):
    """Строки с одинаковой 'Группа' -> одна карточка WB (см. MERGE_BY_GROUP)."""
    if not MERGE_BY_GROUP:
        return [[r] for r in rows]
    grouped = {}
    for r in rows:
        key = r.get("Группа")
        key = f"row{r['_row']}" if key is None else str(key)
        grouped.setdefault(key, []).append(r)
    groups = []
    for key, rs in grouped.items():
        if len(rs) > MAX_VARIANTS_PER_CARD:
            _log(f"!  В группе '{key}' {len(rs)} строк, а в одной карточке WB — не больше "
                  f"{MAX_VARIANTS_PER_CARD}: разобью на несколько карточек.")
        groups.extend(chunks(rs, MAX_VARIANTS_PER_CARD))
    return groups


def vendors_of(group):
    return [r[ARTICLES_COLUMN] for r in group]


def post_items(items):
    r = api("POST", BASE_URL + "/content/v2/cards/upload", json_body=items)
    body = json_or_empty(r)
    if r.status_code != 200 or body.get("error"):
        text = body.get("errorText") or r.text[:300]
        if body.get("additionalErrors"):
            text += f" | {body['additionalErrors']}"
        return False, f"HTTP {r.status_code}: {text}"
    return True, ""


def send_groups(pending, subject_id, charcs, barcodes, unmatched):
    """Отправляет группы на создание. Возвращает {id(группы): текст ошибки} для тех, что WB отклонил сразу."""
    sync_errors = {}
    for batch in chunks(pending, CREATE_BATCH):
        items = [{"subjectID": subject_id,
                  "variants": [build_variant(r, charcs, barcodes[r[ARTICLES_COLUMN]], unmatched) for r in g]}
                 for g in batch]
        ok, message = post_items(items)
        if ok:
            _log(f"   Отправлено на создание")
        elif len(batch) > 1:
            _log(f"!  WB отклонил пачку ({message}). Отправляю по одной, чтобы найти виновника...")
            for g, item in zip(batch, items):
                time.sleep(UPLOAD_PAUSE)
                ok1, message1 = post_items([item])
                if ok1:
                    _log(f"   Отправлено: {', '.join(vendors_of(g))}")
                else:
                    sync_errors[id(g)] = message1
        else:
            sync_errors[id(batch[0])] = message
        time.sleep(UPLOAD_PAUSE)
    if unmatched:
        _log("!  Эти столбцы таблицы не сопоставились с характеристиками предмета и не отправлены: "
              + ", ".join(sorted(unmatched)))
        unmatched.clear()
    return sync_errors


def wait_created(pending, sync_errors, subject_id, since):
    """Ждёт, пока WB создаст карточки. Возвращает {id(группы): причина} для не создавшихся."""
    failed = {id(g): sync_errors[id(g)] for g in pending if id(g) in sync_errors}
    left = [g for g in pending if id(g) not in sync_errors]
    if not left:
        return failed
    _log("Жду, пока WB создаст карточки...")
    deadline = time.time() + WAIT_CREATE_TIMEOUT
    while left:
        time.sleep(WAIT_POLL)
        all_vendors = {v for g in left for v in vendors_of(g)}
        cards = fetch_cards(subject_id, wanted=all_vendors, since=since)
        errors = get_creation_errors(all_vendors, since)
        still = []
        for g in left:
            vs = vendors_of(g)
            if all(v in cards for v in vs):
                continue
            errs = [f"{v}: {errors[v]}" for v in vs if v in errors]
            if errs:
                failed[id(g)] = " | ".join(errs)
                continue
            still.append(g)
        left = still
        if left and time.time() > deadline:
            for g in left:
                failed[id(g)] = f"WB не создал карточку за {WAIT_CREATE_TIMEOUT} сек и не вернул ошибок"
            break
        if left:
            _log(f"   Ещё не создано: {len(left)}...")
    return failed


def create_cards(todo, subject_id, charcs):
    groups = make_groups(todo)
    articles = [r[ARTICLES_COLUMN] for r in todo]
    barcodes = dict(zip(articles, generate_barcodes(len(articles))))
    if len(barcodes) != len(articles):
        raise SystemExit("x WB выдал меньше штрихкодов, чем нужно.")
    unmatched = set()
    pending = groups
    reasons = {}

    for attempt in range(RETRIES + 1):
        if attempt:
            _log(f"\nПовторная попытка {attempt} из {RETRIES}: не создано {len(pending)} карточек")
        since = datetime.now(timezone.utc) - timedelta(minutes=10)
        sync_errors = send_groups(pending, subject_id, charcs, barcodes, unmatched)
        failed = wait_created(pending, sync_errors, subject_id, since)
        if not failed:
            _log("   Все карточки созданы.")
            return
        pending = [g for g in pending if id(g) in failed]
        reasons = {id(g): failed[id(g)] for g in pending}
        _log("\nНе создались:")
        for g in pending:
            _log(f"   {', '.join(vendors_of(g))}\n      причина: {reasons[id(g)]}")
        # перед повтором убеждаемся, что карточки не появились в WB за это время (защита от дублей)
        current = fetch_cards(subject_id, wanted={v for g in pending for v in vendors_of(g)})
        pending = [g for g in pending if not all(v in current for v in vendors_of(g))]
        if not pending:
            _log("   Карточки всё же появились в WB.")
            return

    _log("\n" + "=" * 60)
    _log(f"x Карточки не создались даже после {RETRIES} повторов. Остановка.")
    _log("   ТН ВЭД и фото не тронуты. Уже созданные карточки остались в WB —")
    _log("   после исправления причины просто запустите скрипт снова: он их пропустит.")
    _log("=" * 60)
    sys.exit(1)


# ------------------- ТН ВЭД -------------------
def current_tnved(card, charc_id):
    for c in card.get("characteristics") or []:
        if c.get("id") == charc_id:
            value = c.get("value")
            if isinstance(value, list):
                value = value[0] if value else ""
            return str(value).strip()
    return ""


def card_for_update(card, charc_id, code):
    """cards/update перезаписывает карточку целиком — отправляем все текущие данные + новый ТН ВЭД."""
    characteristics = [{"id": c["id"], "value": c["value"]}
                       for c in (card.get("characteristics") or []) if c.get("id") != charc_id]
    characteristics.append({"id": charc_id, "value": [code]})
    raw_dimensions = card.get("dimensions") or {}
    if isinstance(raw_dimensions, list):
        raw_dimensions = raw_dimensions[0] if raw_dimensions else {}
    dimensions = {
        k: v for k, v in raw_dimensions.items() if k != "isValid"
    } if isinstance(raw_dimensions, dict) else {}
    sizes = [{"chrtID": s["chrtID"], "techSize": s.get("techSize", ""), "wbSize": s.get("wbSize", ""),
              "skus": s.get("skus") or []} for s in (card.get("sizes") or [])]
    return {
        "nmID": card["nmID"],
        "vendorCode": card["vendorCode"],
        "brand": card.get("brand", ""),
        "title": card.get("title", ""),
        "description": card.get("description", ""),
        "dimensions": dimensions,
        "characteristics": characteristics,
        "sizes": sizes,
    }


def step_tnved(rows, cards, subject_id, charcs):
    """Возвращает список артикулов, где ТН ВЭД проставить не удалось."""
    charc_id, charc_name = find_tnved_charc(charcs)
    _log(f"Характеристика ТН ВЭД: '{charc_name}' (ID {charc_id})")

    wanted = {}
    updates = []
    for row in rows:
        article = row[ARTICLES_COLUMN]
        code = tnved_of_row(row)
        wanted[article] = code
        card = cards.get(article)
        if not card:
            continue
        if current_tnved(card, charc_id) == code:
            continue
        updates.append(card_for_update(card, charc_id, code))

    _log(f"ТН ВЭД нужно проставить в {len(updates)} карточках")
    if not updates:
        return [a for a in wanted if a not in cards]

    for batch in chunks(updates, UPDATE_BATCH):
        r = api("POST", BASE_URL + "/content/v2/cards/update", json_body=batch)
        body = json_or_empty(r)
        if r.status_code != 200 or body.get("error"):
            _log(f"!  Ошибка обновления (HTTP {r.status_code}): {body.get('errorText') or r.text[:300]}")
            if body.get("additionalErrors"):
                _log(f"   {body['additionalErrors']}")
        else:
            _log(f"   Отправлено на обновление: {len(batch)} шт.")
        time.sleep(UPLOAD_PAUSE)

    # WB применяет обновление не мгновенно — проверяем результат
    _log("Проверяю, что ТН ВЭД применился...")
    deadline = time.time() + WAIT_TNVED_TIMEOUT
    bad = list(wanted)
    while True:
        fresh = fetch_cards(subject_id, wanted=set(wanted))
        bad = [a for a, code in wanted.items()
               if a not in fresh or current_tnved(fresh[a], charc_id) != code]
        if not bad or time.time() > deadline:
            break
        _log(f"   Ещё не применён: {len(bad)}...")
        time.sleep(15)
    if bad:
        for a in bad:
            _log(f"!  ТН ВЭД не применился: {a}")
    else:
        _log("   ТН ВЭД везде проставлен.")
    return bad


# ------------------- ФОТО -------------------
def upload_photo(nmid, number, path):
    headers = dict(HEADERS)
    headers["X-Nm-Id"] = str(nmid)
    headers["X-Photo-Number"] = str(number)
    for _ in range(5):
        try:
            with open(path, "rb") as f:
                r = requests.post(BASE_URL + "/content/v3/media/file", headers=headers,
                                  files={"uploadfile": (path.name, f, "image/jpeg")}, timeout=120)
            if r.status_code == 200:
                return True
            if r.status_code in (429, 500, 502, 503, 504):
                try:
                    wait = float(r.headers.get("X-Ratelimit-Retry", 0))
                except ValueError:
                    wait = 0
                time.sleep(wait or 3)
                continue
            _log(r.text)
            return False
        except requests.RequestException:
            time.sleep(3)
    return False


def step_photos(rows, cards, folder_index):
    """Возвращает список артикулов, где часть фото не загрузилась."""
    static = collect_static_photos()
    problems = []
    for row in rows:
        article = row[ARTICLES_COLUMN]
        card = cards.get(article)
        if not card:
            _log(f"{article} : карточка не найдена")
            problems.append(article)
            continue
        folder = folder_index.get(row["_orig_article"].lower())
        all_photos = collect_photos(folder) + static
        current = len(card.get("photos") or [])
        if not FORCE_REUPLOAD and current >= len(all_photos):
            _log(f"⏭ {article} : уже есть {current} фото, пропускаем загрузку.")
            continue
        _log(article)
        failed_any = False
        for number, path in enumerate(all_photos, start=1):
            ok = upload_photo(card["nmID"], number, path)
            failed_any = failed_any or not ok
            _log(f"   {'✓' if ok else '✗'} {number} ({path.name})")
            time.sleep(0.5)
        if failed_any:
            problems.append(article)
        time.sleep(1)
    return problems


# ------------------- MAIN -------------------
def main(
    artist=None,
    log_callback=None,
    progress_callback=None,
    prompt_callback=None,
    config_path=None,
):
    """Полный запуск WB-конвейера.

    Для GUI:
        result = main(
            artist="Pink Floyd",
            log_callback=gui_log,
            progress_callback=gui_progress,
            prompt_callback=gui_prompt,
        )

    Для Terminal:
        main()

    Возвращает словарь с результатом, чтобы GUI не приходилось
    разбирать текст консоли.
    """
    global API_TOKEN, HEADERS
    global LOG_CALLBACK, PROMPT_CALLBACK, PROGRESS_CALLBACK

    LOG_CALLBACK = log_callback
    PROMPT_CALLBACK = prompt_callback
    PROGRESS_CALLBACK = progress_callback

    result = {
        "success": False,
        "artist": None,
        "rows": 0,
        "existing_cards": 0,
        "created_candidates": 0,
        "tnved_errors": [],
        "photo_errors": [],
        "error": None,
    }

    try:
        _progress(0, 100, "Запуск")
        API_TOKEN = load_config(config_path)
        HEADERS = {"Authorization": API_TOKEN}

        if not EXCEL_FILE.exists():
            raise RuntimeError(f"Не найден файл '{EXCEL_FILE}'")

        _log("=" * 60)
        _log("WB UPLOADER: карточки -> ТН ВЭД -> фото")
        _log("=" * 60)

        artist = get_artist_name(artist)
        result["artist"] = artist
        _progress(5, 100, "Проверка файлов")

        artist_folder = find_artist_output_folder(artist)
        if not artist_folder:
            raise RuntimeError(
                f"Не найдена папка '{OUTPUT_FOLDER / artist}' — "
                "сначала запустите mockuper.jsx."
            )
        folder_index = build_folder_index(artist_folder, artist)

        rows = read_table()
        result["rows"] = len(rows)
        if not rows:
            raise RuntimeError("В таблице нет строк с артикулами.")
        _log(f"Строк в таблице: {len(rows)}")

        # 0. Локальные проверки до любых запросов к WB.
        _progress(10, 100, "Локальная проверка Excel и фотографий")
        validate_group_limits(rows)
        validate_local(rows, folder_index)

        categories = {r.get("Категория продавца") for r in rows}
        if len(categories) != 1 or None in categories:
            raise RuntimeError(
                "В столбце 'Категория продавца' должно быть одно значение, "
                f"сейчас: {categories}"
            )
        subject_name = categories.pop()

        # Справочники WB.
        _progress(15, 100, "Получение справочников WB")
        subject_id = find_subject_id(subject_name)
        charcs = load_charcs(subject_id)
        _log(
            f"Предмет '{subject_name}' (subjectID {subject_id}), "
            f"характеристик: {len(charcs)}"
        )

        check_color_values(rows)
        check_tnved_codes(rows, subject_id)

        # Список существующих карточек.
        _progress(25, 100, "Проверка существующих карточек WB")
        _log("Загружаю список карточек из WB...")
        existing = fetch_cards(subject_id)
        result["existing_cards"] = sum(
            1 for r in rows if r[ARTICLES_COLUMN] in existing
        )
        _log(f"   Карточек в WB по этому предмету: {len(existing)}")

        todo = [
            r for r in rows
            if r[ARTICLES_COLUMN] not in existing
        ]
        _log(
            f"В таблице: {len(rows)}, уже есть в WB: "
            f"{len(rows) - len(todo)}, к созданию: {len(todo)}"
        )

        # 1. Правка длинных значений.
        renames = load_renames()
        if todo:
            _progress(30, 100, "Проверка длинных названий и артикулов")
            fix_lengths(todo, rows, renames)

            todo = [
                r for r in todo
                if r[ARTICLES_COLUMN] not in existing
            ]

            if MERGE_BY_GROUP:
                partial = {}
                for g in make_groups(todo):
                    key = str(g[0].get("Группа"))
                    if any(
                        str(r.get("Группа")) == key
                        and r[ARTICLES_COLUMN] in existing
                        for r in rows
                    ):
                        partial[key] = vendors_of(g)

                if partial:
                    _log(
                        "\n! Часть строк из этих групп уже есть в WB. "
                        "Недостающие будут созданы отдельными карточками "
                        "и не объединятся с существующими:"
                    )
                    for key, vs in partial.items():
                        _log(f"   группа {key}: {', '.join(vs)}")

                    if not _confirm("   Продолжить? [y/N]: "):
                        raise RuntimeError("Остановлено пользователем.")

        result["created_candidates"] = len(todo)

        # 2. Создание карточек.
        if todo:
            _progress(35, 100, f"Создание карточек: {len(todo)}")
            _log("\n--- Создание карточек ---")
            create_cards(todo, subject_id, charcs)

        articles = {r[ARTICLES_COLUMN] for r in rows}
        _progress(55, 100, "Обновление списка карточек")
        cards = fetch_cards(subject_id, wanted=articles)

        # 3. ТН ВЭД.
        _progress(60, 100, "Проверка и установка ТН ВЭД")
        _log("\n--- ТН ВЭД ---")
        tnved_bad = step_tnved(rows, cards, subject_id, charcs)
        result["tnved_errors"] = list(tnved_bad)

        # 4. Фото.
        _progress(75, 100, "Загрузка фотографий")
        _log("\n--- Фото ---")
        cards = fetch_cards(subject_id, wanted=articles)
        photo_bad = step_photos(rows, cards, folder_index)
        result["photo_errors"] = list(photo_bad)

        _progress(100, 100, "Завершено")
        _log("\n" + "=" * 60)

        if tnved_bad or photo_bad:
            _log("ГОТОВО, НО ЕСТЬ ПРОБЛЕМЫ:")
            if tnved_bad:
                _log("  ТН ВЭД не проставлен:")
                for a in tnved_bad:
                    _log(f"    {a}")
            if photo_bad:
                _log("  Фото загружены не полностью:")
                for a in photo_bad:
                    _log(f"    {a}")
            _log("Запустите скрипт ещё раз — выполненное он пропустит.")
            result["success"] = False
        else:
            _log(
                "ГОТОВО! Карточки созданы, ТН ВЭД и фото проставлены. "
                "Осталось выставить остатки."
            )
            result["success"] = True

            try:
                if RENAMES_FILE.exists():
                    RENAMES_FILE.unlink()
            except OSError:
                pass

        _log("=" * 60)
        return result

    except SystemExit as e:
        message = str(e) or "Операция остановлена."
        result["error"] = message
        _log(f"\nОШИБКА: {message}")
        return result

    except Exception as e:
        result["error"] = str(e)
        _log(f"\nОШИБКА: {e}")
        return result




if __name__ == "__main__":
    result = main()
    if result["success"]:
        _log("\nРезультат: УСПЕШНО")
    else:
        _log(f"\nРезультат: ОШИБКА — {result['error']}")